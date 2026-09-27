"""
export_utils.py
----------------
Sinh bao cao xuat ra 3 dinh dang: Excel (.xlsx), Word (.docx), PDF (.pdf).
Du lieu lay truc tiep tu MongoDB (qua queries.py) tai thoi diem export,
nen bao cao luon phan anh dung du lieu moi nhat.
"""

import os
import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import BarChart, PieChart, Reference

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
)
from reportlab.lib.styles import getSampleStyleSheet

import queries as q

EXPORT_DIR = os.path.join(os.path.dirname(__file__), "..", "exports")
os.makedirs(EXPORT_DIR, exist_ok=True)


def _timestamp():
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")


# --------------------------------------------------------------------------
# EXCEL
# --------------------------------------------------------------------------
def export_excel():
    wb = Workbook()

    # ----- Sheet 1: KPI tong quan -----
    ws1 = wb.active
    ws1.title = "Tong quan"
    kpi = q.kpi_summary()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")

    ws1["A1"] = "BAO CAO BI - TUYEN SINH NEU/PTIT"
    ws1["A1"].font = Font(bold=True, size=14)
    ws1["A2"] = f"Xuat luc: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}"

    row = 4
    ws1[f"A{row}"] = "Chi so"
    ws1[f"B{row}"] = "Gia tri"
    ws1[f"A{row}"].font = header_font
    ws1[f"B{row}"].font = header_font
    ws1[f"A{row}"].fill = header_fill
    ws1[f"B{row}"].fill = header_fill
    labels = {
        "total_contacts": "Tong so lien he",
        "paid_contacts": "Da thanh toan",
        "unpaid_contacts": "Chua thanh toan",
        "total_revenue": "Tong doanh thu (VND)",
        "need_check_count": "So can kiem tra lai",
        "web_registered_count": "Dang ky qua website",
    }
    row += 1
    for key, label in labels.items():
        ws1[f"A{row}"] = label
        ws1[f"B{row}"] = kpi.get(key, 0)
        row += 1
    ws1.column_dimensions["A"].width = 28
    ws1.column_dimensions["B"].width = 20

    # ----- Sheet 2: theo khu vuc -----
    ws2 = wb.create_sheet("Theo khu vuc")
    ws2.append(["Khu vuc", "So luong"])
    for c in ws2[1]:
        c.font = header_font
        c.fill = header_fill
    for row_data in q.group_count("region"):
        ws2.append([row_data["label"], row_data["count"]])
    chart = PieChart()
    chart.title = "Phan bo theo khu vuc"
    data_ref = Reference(ws2, min_col=2, min_row=1, max_row=ws2.max_row)
    cat_ref = Reference(ws2, min_col=1, min_row=2, max_row=ws2.max_row)
    chart.add_data(data_ref, titles_from_data=True)
    chart.set_categories(cat_ref)
    ws2.add_chart(chart, "D2")

    # ----- Sheet 3: theo truong hoc -----
    ws3 = wb.create_sheet("Theo truong")
    ws3.append(["Truong", "So luong"])
    for c in ws3[1]:
        c.font = header_font
        c.fill = header_fill
    for row_data in q.group_count("school"):
        ws3.append([row_data["label"], row_data["count"]])
    chart3 = BarChart()
    chart3.title = "So luong theo truong"
    data_ref = Reference(ws3, min_col=2, min_row=1, max_row=ws3.max_row)
    cat_ref = Reference(ws3, min_col=1, min_row=2, max_row=ws3.max_row)
    chart3.add_data(data_ref, titles_from_data=True)
    chart3.set_categories(cat_ref)
    ws3.add_chart(chart3, "D2")

    # ----- Sheet 4: trang thai thanh toan -----
    ws4 = wb.create_sheet("Thanh toan")
    ws4.append(["Trang thai", "So luong"])
    for c in ws4[1]:
        c.font = header_font
        c.fill = header_fill
    for row_data in q.group_count("payment_status"):
        ws4.append([row_data["label"], row_data["count"]])

    # ----- Sheet 5: danh sach chi tiet (raw data cho pivot) -----
    ws5 = wb.create_sheet("Chi tiet lien he")
    contacts = q.all_contacts_for_export()
    if contacts:
        cols = list(contacts[0].keys())
        ws5.append(cols)
        for c in ws5[1]:
            c.font = header_font
            c.fill = header_fill
        for c in contacts:
            ws5.append([c.get(col, "") for col in cols])
        for i, col in enumerate(cols, start=1):
            ws5.column_dimensions[ws5.cell(row=1, column=i).column_letter].width = 18

    fname = f"BaoCao_BI_{_timestamp()}.xlsx"
    path = os.path.join(EXPORT_DIR, fname)
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# WORD
# --------------------------------------------------------------------------
def export_word():
    doc = Document()

    title = doc.add_heading("BAO CAO BI - TUYEN SINH NEU/PTIT", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph(f"Xuat luc: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_heading("1. Chi so tong quan", level=1)
    kpi = q.kpi_summary()
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    hdr[0].text = "Chi so"
    hdr[1].text = "Gia tri"
    labels = {
        "total_contacts": "Tong so lien he",
        "paid_contacts": "Da thanh toan",
        "unpaid_contacts": "Chua thanh toan",
        "total_revenue": "Tong doanh thu (VND)",
        "need_check_count": "So can kiem tra lai",
        "web_registered_count": "Dang ky qua website",
    }
    for key, label in labels.items():
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(kpi.get(key, 0))

    doc.add_heading("2. Phan bo theo khu vuc", level=1)
    t2 = doc.add_table(rows=1, cols=2)
    t2.style = "Light Grid Accent 1"
    t2.rows[0].cells[0].text = "Khu vuc"
    t2.rows[0].cells[1].text = "So luong"
    for r in q.group_count("region"):
        c = t2.add_row().cells
        c[0].text = str(r["label"])
        c[1].text = str(r["count"])

    doc.add_heading("3. Phan bo theo truong", level=1)
    t3 = doc.add_table(rows=1, cols=2)
    t3.style = "Light Grid Accent 1"
    t3.rows[0].cells[0].text = "Truong"
    t3.rows[0].cells[1].text = "So luong"
    for r in q.group_count("school"):
        c = t3.add_row().cells
        c[0].text = str(r["label"])
        c[1].text = str(r["count"])

    doc.add_heading("4. Trang thai thanh toan", level=1)
    t4 = doc.add_table(rows=1, cols=2)
    t4.style = "Light Grid Accent 1"
    t4.rows[0].cells[0].text = "Trang thai"
    t4.rows[0].cells[1].text = "So luong"
    for r in q.group_count("payment_status"):
        c = t4.add_row().cells
        c[0].text = str(r["label"])
        c[1].text = str(r["count"])

    doc.add_heading("5. Drip / Segment marketing", level=1)
    t5 = doc.add_table(rows=1, cols=3)
    t5.style = "Light Grid Accent 1"
    t5.rows[0].cells[0].text = "Drip"
    t5.rows[0].cells[1].text = "Segment"
    t5.rows[0].cells[2].text = "So lien he"
    for r in q.drip_segments_overview():
        c = t5.add_row().cells
        c[0].text = str(r.get("drip", ""))
        c[1].text = str(r.get("segment", ""))
        c[2].text = str(r.get("contact_count", ""))

    fname = f"BaoCao_BI_{_timestamp()}.docx"
    path = os.path.join(EXPORT_DIR, fname)
    doc.save(path)
    return path


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def export_pdf():
    fname = f"BaoCao_BI_{_timestamp()}.pdf"
    path = os.path.join(EXPORT_DIR, fname)

    doc = SimpleDocTemplate(path, pagesize=A4,
                             leftMargin=2 * cm, rightMargin=2 * cm,
                             topMargin=1.5 * cm, bottomMargin=1.5 * cm)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("BAO CAO BI - TUYEN SINH NEU/PTIT", styles["Title"]))
    elements.append(Paragraph(
        f"Xuat luc: {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')}", styles["Normal"]))
    elements.append(Spacer(1, 0.5 * cm))

    def add_table(title, header, rows):
        elements.append(Paragraph(title, styles["Heading2"]))
        data = [header] + rows
        t = Table(data, hAlign="LEFT")
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F5597")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 0.5 * cm))

    kpi = q.kpi_summary()
    labels = {
        "total_contacts": "Tong so lien he",
        "paid_contacts": "Da thanh toan",
        "unpaid_contacts": "Chua thanh toan",
        "total_revenue": "Tong doanh thu (VND)",
        "need_check_count": "So can kiem tra lai",
        "web_registered_count": "Dang ky qua website",
    }
    add_table("1. Chi so tong quan", ["Chi so", "Gia tri"],
              [[label, str(kpi.get(key, 0))] for key, label in labels.items()])

    add_table("2. Phan bo theo khu vuc", ["Khu vuc", "So luong"],
              [[str(r["label"]), str(r["count"])] for r in q.group_count("region")])

    add_table("3. Phan bo theo truong", ["Truong", "So luong"],
              [[str(r["label"]), str(r["count"])] for r in q.group_count("school")])

    add_table("4. Trang thai thanh toan", ["Trang thai", "So luong"],
              [[str(r["label"]), str(r["count"])] for r in q.group_count("payment_status")])

    doc.build(elements)
    return path
